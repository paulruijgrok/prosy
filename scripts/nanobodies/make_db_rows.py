#!/usr/bin/env python3
"""Turn a nanobody plate run into database rows for the lab sheet.

Reads the provenance mapping CSV written by ``make_nanobody_plate.py`` (which
carries the *new* sequential IDs, e.g. ``Nb73``) plus the lab spreadsheet the
parents came from, and writes an xlsx whose columns line up with the
``Nanobodies`` sheet of ``Malaria DX plasmids.xlsx`` — so the block can be
copy-pasted straight underneath the existing rows.

Columns that the lab sheet fills with its own formulas (the long name, the
length) are deliberately left **empty**: pasting values there would clobber the
formulas. Everything else is filled from the run.

Example
-------
    python make_db_rows.py --stamp 260626_row \
        --data-dir "../../Working folder/260720_NanobodyMuts_row_gc72"
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from pathlib import Path

# Header cells of the `Nanobodies` sheet that the sheet computes itself.
# We emit the column (so the block stays aligned) but leave every cell blank.
FORMULA_COLUMNS = ("Long", "Length (AA)")

# Sub-header row of the lab sheet: row 1 spans "Nanobody ID" over two columns,
# row 2 names them. Merge them so column keys are unique and readable.
SUBHEADER_ROW = 2

CDR_COLUMNS = ("CDR 1", "CDR 2", "CDR 3")

# --include choice -> the mapping CSV's `Type` values it selects.
INCLUDE_TYPES = {"mutants": {"mutant"}, "parents": {"parent"},
                 "all": {"mutant", "parent"}}


# --------------------------------------------------------------------------- #
# Reading                                                                     #
# --------------------------------------------------------------------------- #


@dataclass
class SheetInfo:
    """What we need to know about the existing `Nanobodies` sheet."""

    columns: list[str]                 # output column order, sub-headers merged
    by_id: dict[str, dict]             # NbID -> row dict
    next_db_id: int                    # first unused DB ID
    first_empty_row: int               # 1-based worksheet row to paste into


def read_sheet(path: Path, sheet: str, id_col: str, db_id_col: str) -> SheetInfo:
    """Read the lab sheet: column order, existing rows, and where to paste."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < SUBHEADER_ROW:
        raise SystemExit(f"Sheet {sheet!r} in {path} has no header rows.")

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    sub = [str(c).strip() if c is not None else "" for c in rows[SUBHEADER_ROW - 1]]
    columns = [sub[i] if not header[i] and i < len(sub) else header[i]
               for i in range(len(header))]
    # Trim trailing unnamed columns; they carry no meaning for the paste block.
    while columns and not columns[-1]:
        columns.pop()
    if id_col not in columns:
        raise SystemExit(f"Column {id_col!r} not found in sheet {sheet!r}.")

    by_id: dict[str, dict] = {}
    max_db_id = 0
    first_empty_row = SUBHEADER_ROW + 1
    for r_idx, row in enumerate(rows[SUBHEADER_ROW:], start=SUBHEADER_ROW + 1):
        rec = {columns[i]: (row[i] if i < len(row) else None)
               for i in range(len(columns))}
        nb_id = rec.get(id_col)
        if not nb_id:
            continue
        by_id[str(nb_id).strip()] = rec
        first_empty_row = r_idx + 1
        db_id = rec.get(db_id_col)
        if isinstance(db_id, (int, float)):
            max_db_id = max(max_db_id, int(db_id))

    return SheetInfo(columns=columns, by_id=by_id,
                     next_db_id=max_db_id + 1, first_empty_row=first_empty_row)


def read_mapping(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        raise SystemExit(f"Mapping file is empty: {path}")
    required = {"Well", "NbID", "NbID_parent", "Parent", "Set", "Mutation",
                "Score", "Type", "Protein"}
    missing = required - set(rows[0])
    if missing:
        raise SystemExit(f"Mapping {path} is missing columns: {sorted(missing)}")
    return rows


# --------------------------------------------------------------------------- #
# CDRs                                                                        #
# --------------------------------------------------------------------------- #


def cdrs_at_parent_offsets(parent_protein: str, parent_cdrs: dict[str, str],
                           protein: str) -> tuple[dict[str, str], list[str]]:
    """Re-extract each CDR from ``protein`` at the parent's CDR offsets.

    Point mutations preserve length, so the parent's CDR boundaries transfer
    directly. A mutation landing inside a CDR therefore shows up in the result.
    Returns the CDRs plus any warnings (CDR absent from, or ambiguous in, the
    parent sequence — those are copied verbatim instead).
    """
    out: dict[str, str] = {}
    warnings: list[str] = []
    for col in CDR_COLUMNS:
        cdr = (parent_cdrs.get(col) or "").strip()
        if not cdr:
            out[col] = ""
            continue
        start = parent_protein.find(cdr)
        if start < 0:
            warnings.append(f"{col} {cdr!r} not found in parent sequence")
            out[col] = cdr
        elif parent_protein.find(cdr, start + 1) >= 0:
            warnings.append(f"{col} {cdr!r} occurs more than once in parent")
            out[col] = cdr
        elif len(protein) != len(parent_protein):
            warnings.append(f"{col}: variant length differs from parent")
            out[col] = cdr
        else:
            out[col] = protein[start:start + len(cdr)]
    return out, warnings


# --------------------------------------------------------------------------- #
# Row building                                                                #
# --------------------------------------------------------------------------- #


def build_rows(mapping: list[dict], info: SheetInfo, *, include: str,
               cdr_mode: str, db_id_start: int, id_col: str, db_id_col: str,
               stamp: str) -> tuple[list[dict], list[str]]:
    """Build one output dict per selected variant, in plate order."""
    warnings: list[str] = []
    wanted = INCLUDE_TYPES[include]
    selected = [r for r in mapping if r["Type"] in wanted]
    rows: list[dict] = []
    db_id = db_id_start

    for rec in selected:
        parent_id = rec["Parent"].strip()
        parent = info.by_id.get(parent_id)
        if parent is None:
            raise SystemExit(
                f"Parent {parent_id!r} not found in the lab sheet — cannot "
                f"inherit its metadata for {rec['NbID']}.")
        protein = rec["Protein"].strip()
        is_parent = rec["Type"] == "parent"

        if cdr_mode == "blank":
            cdrs = {c: "" for c in CDR_COLUMNS}
        elif cdr_mode == "inherit" or is_parent:
            cdrs = {c: (parent.get(c) or "") for c in CDR_COLUMNS}
        else:
            parent_protein = str(parent.get("AA Sequence") or "").strip()
            cdrs, warns = cdrs_at_parent_offsets(
                parent_protein, {c: parent.get(c) for c in CDR_COLUMNS}, protein)
            warnings += [f"{rec['NbID']} ({parent_id}): {w}" for w in warns]

        parent_alias = str(parent.get("Alias") or "").strip()
        mutation = rec["Mutation"].strip()
        alias = (f"{parent_alias}_{mutation}"
                 if parent_alias and mutation else parent_alias)
        note = (f"Mutant of {parent_id} ({mutation}); design score "
                f"{rec['Score']}; plate {stamp} well {rec['Well']}"
                if not is_parent else f"Plate {stamp} well {rec['Well']}")

        row = {c: "" for c in info.columns}
        row.update({
            db_id_col: db_id,
            id_col: rec["NbID"],
            "Target": parent.get("Target") or "",
            "Target spec": parent.get("Target spec") or "",
            "Source": parent.get("Source") or "",
            "Set name": rec["Set"] or (parent.get("Set name") or ""),
            "AA Sequence": protein,
            "Alias": alias,
            "Parent": parent_id,
            "Notes": note,
            **cdrs,
        })
        for col in FORMULA_COLUMNS:
            if col in row:
                row[col] = ""
        rows.append(row)
        db_id += 1

    return rows, warnings


def write_xlsx(rows: list[dict], columns: list[str], path: Path) -> Path:
    """Write the paste block: one header row, then the data rows."""
    import openpyxl

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "db_rows"
    ws.append(columns)
    for r in rows:
        ws.append([r.get(c, "") for c in columns])
    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
# CLI                                                                         #
# --------------------------------------------------------------------------- #


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    here = Path(__file__).resolve().parent
    default_data = here / ".." / ".." / "Working folder" / "260626_NanobodyMuts"

    p = argparse.ArgumentParser(
        description="Build lab-sheet database rows from a nanobody plate run.")
    p.add_argument("--data-dir", type=Path, default=default_data,
                   help="Folder holding the mapping CSV and the lab xlsx.")
    p.add_argument("--stamp", default="260626",
                   help="Filename prefix of the run (as passed to make_nanobody_plate.py).")
    p.add_argument("--mapping", type=Path, default=None,
                   help="Mapping CSV (default: <data-dir>/<stamp>_nanobody_mapping.csv).")
    p.add_argument("--sheet-xlsx", type=Path, default=None,
                   help="Lab spreadsheet (default: <data-dir>/Malaria DX plasmids.xlsx).")
    p.add_argument("--sheet", default="Nanobodies", help="Worksheet to mirror.")
    p.add_argument("--out", type=Path, default=None,
                   help="Output xlsx (default: <data-dir>/<stamp>_nanobody_db_rows.xlsx).")
    p.add_argument("--include", choices=["mutants", "parents", "all"],
                   default="mutants",
                   help="Which variants to emit (parents already exist in the sheet).")
    p.add_argument("--cdr-mode", choices=["offsets", "inherit", "blank"],
                   default="offsets",
                   help="offsets: re-extract CDRs at the parent's positions "
                        "(mutations inside a CDR show up); inherit: copy the "
                        "parent's CDRs verbatim; blank: leave empty.")
    p.add_argument("--db-id-start", type=int, default=None,
                   help="First DB ID (default: highest in the sheet + 1).")
    p.add_argument("--id-col", default="Nanobody ID", help="Short-ID column name.")
    p.add_argument("--db-id-col", default="DB ID", help="DB ID column name.")
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    data_dir = args.data_dir.resolve()
    mapping_path = args.mapping or data_dir / f"{args.stamp}_nanobody_mapping.csv"
    sheet_path = args.sheet_xlsx or data_dir / "Malaria DX plasmids.xlsx"
    out_path = args.out or data_dir / f"{args.stamp}_nanobody_db_rows.xlsx"

    for path, what in ((mapping_path, "mapping CSV"), (sheet_path, "lab spreadsheet")):
        if not path.exists():
            raise SystemExit(f"{what} not found: {path}")

    info = read_sheet(sheet_path, args.sheet, args.id_col, args.db_id_col)
    mapping = read_mapping(mapping_path)
    db_id_start = args.db_id_start if args.db_id_start is not None else info.next_db_id

    rows, warnings = build_rows(
        mapping, info, include=args.include, cdr_mode=args.cdr_mode,
        db_id_start=db_id_start, id_col=args.id_col, db_id_col=args.db_id_col,
        stamp=args.stamp)
    if not rows:
        raise SystemExit(f"No {args.include} rows in {mapping_path}.")

    clashes = sorted({r[args.id_col] for r in rows} & set(info.by_id))
    if clashes:
        warnings.append(f"IDs already present in the sheet: {', '.join(clashes)}")

    write_xlsx(rows, info.columns, out_path)

    ids = [r[args.id_col] for r in rows]
    print(f"Wrote {len(rows)} {args.include} rows ({ids[0]}–{ids[-1]}) to {out_path}")
    print(f"  columns : {len(info.columns)} (mirroring sheet {args.sheet!r}; "
          f"{', '.join(FORMULA_COLUMNS)} left blank for the sheet's formulas)")
    print(f"  DB ID   : {db_id_start}–{db_id_start + len(rows) - 1}")
    print(f"  paste at: row {info.first_empty_row} of {args.sheet!r} "
          f"(data rows only — skip the header row of the output file)")
    for w in warnings:
        print(f"  WARNING : {w}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
