"""Reading and writing the tabular formats that show up in synthesis workflows:
CSV/TSV design files, the lab's xlsx database, and vendor plate-upload sheets.

Flexible, case-insensitive column matching mirrors the original script so messy
spreadsheets still resolve to the expected fields.
"""

from __future__ import annotations

import csv
import shutil
import subprocess
from pathlib import Path

# Plate-upload column headers (as in example_plate-file-upload.xls).
PLATE_COLUMNS = ["Well Position", "Name", "Sequence"]
# Design-file columns (as in design_seqs_miniprotein.csv).
DESIGN_COLUMNS = ["Name", "Sequence", "5' nucleotides", "3' nucleotides"]


def _strip_bom(s: str) -> str:
    return s.lstrip("﻿").strip() if isinstance(s, str) else s


def find_columns(available: list[str], required: dict[str, str]) -> dict[str, str]:
    """Map logical keys to actual column names, case-insensitively.

    ``required`` maps a logical key -> expected header. Raises ValueError if a
    required header cannot be found.
    """
    norm = {_strip_bom(c).lower(): c for c in available}
    out: dict[str, str] = {}
    for key, header in required.items():
        actual = norm.get(header.lower())
        if actual is None:
            raise ValueError(
                f"Required column {header!r} not found. Available: {available}"
            )
        out[key] = actual
    return out


def read_delimited(path: str | Path, delimiter: str | None = None) -> list[dict]:
    """Read a CSV/TSV into a list of dicts (BOM-safe, delimiter auto-detected)."""
    path = Path(path)
    if delimiter is None:
        delimiter = "\t" if path.suffix.lower() in {".tsv", ".tab"} else ","
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        rows = [{(_strip_bom(k)): v for k, v in row.items()} for row in reader]
    return rows


def read_xlsx_sheet(path: str | Path, sheet: str | None = None) -> list[dict]:
    """Read a worksheet into a list of dicts using the first row as headers."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [(_strip_bom(c) if c is not None else "") for c in next(rows_iter)]
    except StopIteration:
        return []
    out = []
    for row in rows_iter:
        out.append({header[i]: row[i] if i < len(row) else None for i in range(len(header))})
    return out


def write_csv(rows: list[dict], path: str | Path, columns: list[str]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for r in rows:
            writer.writerow(r)


def write_plate_xlsx(
    records: list[dict],
    path: str | Path,
    *,
    name_key: str = "name",
    sequence_key: str = "sequence",
    well_key: str = "well",
) -> Path:
    """Write a vendor-style plate sheet (Well Position / Name / Sequence)."""
    import openpyxl

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plate"
    ws.append(PLATE_COLUMNS)
    for rec in records:
        ws.append([rec[well_key], rec[name_key], rec[sequence_key]])
    wb.save(path)
    return path


def convert_to_xls(xlsx_path: str | Path) -> Path | None:
    """Best-effort convert .xlsx -> legacy .xls via LibreOffice if available.

    Returns the .xls path on success, else None (caller keeps the .xlsx).
    """
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice:
        return None
    xlsx_path = Path(xlsx_path)
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "xls", "--outdir",
             str(xlsx_path.parent), str(xlsx_path)],
            check=True, capture_output=True, timeout=120,
        )
    except (subprocess.SubprocessError, OSError):
        return None
    xls = xlsx_path.with_suffix(".xls")
    return xls if xls.exists() else None
